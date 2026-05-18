import random
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date, timedelta
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from apps.api_connector.views import fetch_kpis


def generate_sparkline():
    points = []
    y = random.randint(10, 30)
    for x in range(0, 210, 25):
        y = max(5, min(38, y + random.randint(-8, 8)))
        points.append(f"{x},{y}")
    return " ".join(points)


def get_status(value, target=1.0):
    if value < 0.5:
        return 'good'
    elif value <= 1.0:
        return 'warn'
    else:
        return 'bad'


KPI_TARGETS = {
    'DVX':    1.0,
    'DVT':    1.0,
    'DVM':    1.0,
    'SCA':    1.0,
    'DRR':    1.0,
    'BDACOM': 1.0,
}


def get_zone1_kpis():
    raw_kpis = fetch_kpis('QCP')
    kpis = []
    for raw in raw_kpis:
        name   = raw.get('kpi_type', 'UNKNOWN')
        value  = raw.get('value', 0)
        target = KPI_TARGETS.get(name, 1.0)
        prev   = round(random.uniform(value * 0.85, value * 1.15), 3)
        diff   = round(value - prev, 3)
        status = get_status(value, target)
        kpis.append({
            'name':      name,
            'value':     value,
            'target':    target,
            'status':    status,
            'trend':     'up' if diff > 0 else 'down',
            'change':    f"+{diff}" if diff > 0 else str(diff),
            'min':       round(value * 0.85, 3),
            'max':       round(value * 1.15, 3),
            'sparkline': generate_sparkline(),
        })
    return kpis


@login_required
def home(request):
    zone1_kpis = get_zone1_kpis()
    zone1_good = sum(1 for k in zone1_kpis if k['status'] == 'good')
    zone1_warn = sum(1 for k in zone1_kpis if k['status'] == 'warn')
    zone1_bad  = sum(1 for k in zone1_kpis if k['status'] == 'bad')
    return render(request, 'dashboard/home.html', {
        'zone1_kpis':  zone1_kpis,
        'zone1_good':  zone1_good,
        'zone1_warn':  zone1_warn,
        'zone1_bad':   zone1_bad,
        'periode':     date.today().strftime('%d/%m/%Y'),
        'last_update': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
    })


@login_required
def index(request):
    return redirect('dashboard:home')


@login_required
def zone1(request):
    kpis = get_zone1_kpis()
    good = sum(1 for k in kpis if k['status'] == 'good')
    warn = sum(1 for k in kpis if k['status'] == 'warn')
    bad  = sum(1 for k in kpis if k['status'] == 'bad')
    return render(request, 'dashboard/zone1.html', {
        'kpis': kpis,
        'periode': date.today().strftime('%d/%m/%Y'),
        'last_update': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'summary': {'good': good, 'warn': warn, 'bad': bad},
    })


@login_required
def zone2(request):
    return render(request, 'dashboard/zone2.html', {
        'periode': date.today().strftime('%d/%m/%Y'),
        'last_update': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
    })


@login_required
def zone3(request):
    return render(request, 'dashboard/zone3.html', {
        'periode': date.today().strftime('%d/%m/%Y'),
        'last_update': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
    })


@login_required
def zone1_kpis_partial(request):
    kpis = get_zone1_kpis()
    good = sum(1 for k in kpis if k['status'] == 'good')
    warn = sum(1 for k in kpis if k['status'] == 'warn')
    bad  = sum(1 for k in kpis if k['status'] == 'bad')
    return render(request, 'dashboard/partials/zone1_kpis.html', {
        'kpis': kpis,
        'last_update': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'summary': {'good': good, 'warn': warn, 'bad': bad},
    })


@login_required
def kpi_history(request, kpi_name):
    date_from = request.GET.get('date_from')
    date_to   = request.GET.get('date_to')
    try:
        date_from = datetime.strptime(date_from, '%Y-%m-%d')
        date_to   = datetime.strptime(date_to,   '%Y-%m-%d')
    except (ValueError, TypeError):
        date_to   = datetime.now()
        date_from = date_to - timedelta(days=30)
    target = KPI_TARGETS.get(kpi_name.upper(), 1.0)
    labels = []
    values = []
    current = date_from
    while current <= date_to:
        labels.append(current.strftime('%d/%m'))
        values.append(round(random.uniform(target * 0.6, target * 1.5), 3))
        current += timedelta(days=1)
    return JsonResponse({
        'kpi_name':  kpi_name.upper(),
        'target':    target,
        'labels':    labels,
        'values':    values,
        'date_from': date_from.strftime('%Y-%m-%d'),
        'date_to':   date_to.strftime('%Y-%m-%d'),
    })


@login_required
def export_csv(request):
    kpis = get_zone1_kpis()
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="KPI_QCP_{date.today()}.csv"'
    response.write('\ufeff')
    writer = csv.writer(response, delimiter=';')
    writer.writerow(['Rapport KPI — Zone QCP'])
    writer.writerow([f'Date extraction: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'])
    writer.writerow([f'Période: {date.today().strftime("%d/%m/%Y")}'])
    writer.writerow([])
    writer.writerow(['KPI', 'Valeur', 'Cible', 'Statut', 'Tendance', 'Min', 'Max'])
    for kpi in kpis:
        status_label = {'good': 'Dans la cible', 'warn': 'Alerte', 'bad': 'Hors cible'}.get(kpi['status'], '')
        writer.writerow([
            kpi['name'],
            str(kpi['value']).replace('.', ','),
            str(kpi['target']).replace('.', ','),
            status_label,
            kpi['change'],
            str(kpi['min']).replace('.', ','),
            str(kpi['max']).replace('.', ','),
        ])
    return response


@login_required
def export_excel(request):
    kpis = get_zone1_kpis()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KPI Zone QCP"
    navy   = "1A2C6B"
    green  = "16A34A"
    orange = "D97706"
    red    = "DC2626"
    white  = "FFFFFF"
    gray   = "F9FAFB"
    ws.merge_cells('A1:G1')
    ws['A1'] = 'STELLANTIS — Rapport KPI Zone QCP'
    ws['A1'].font      = Font(name='Calibri', size=14, bold=True, color=white)
    ws['A1'].fill      = PatternFill("solid", fgColor=navy)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    ws['A2'] = 'Date extraction:'
    ws['B2'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
    ws['A3'] = 'Période:'
    ws['B3'] = date.today().strftime('%d/%m/%Y')
    ws['A4'] = 'RSP:'
    ws['B4'] = 'QCP'
    for row in [2, 3, 4]:
        ws.cell(row=row, column=1).font = Font(bold=True, color=navy)
    headers = ['KPI', 'Valeur', 'Cible', 'Statut', 'Tendance', 'Minimum', 'Maximum']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=header)
        cell.font      = Font(name='Calibri', bold=True, color=white, size=11)
        cell.fill      = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[6].height = 22
    status_colors = {'good': green, 'warn': orange, 'bad': red}
    status_labels = {'good': 'Dans la cible', 'warn': 'Alerte', 'bad': 'Hors cible'}
    for row_idx, kpi in enumerate(kpis, 7):
        row_fill = gray if row_idx % 2 == 0 else white
        color    = status_colors.get(kpi['status'], '000000')
        label    = status_labels.get(kpi['status'], '')
        data = [kpi['name'], kpi['value'], kpi['target'], label, kpi['change'], kpi['min'], kpi['max']]
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = Font(name='Calibri', size=10)
            cell.fill      = PatternFill("solid", fgColor=row_fill)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border    = Border(
                bottom=Side(style='thin', color='E5E7EB'),
                right=Side(style='thin',  color='E5E7EB')
            )
            if col_idx == 1:
                cell.font = Font(name='Calibri', size=10, bold=True, color=navy)
            if col_idx == 4:
                cell.font = Font(name='Calibri', size=10, bold=True, color=color)
        ws.row_dimensions[row_idx].height = 20
    widths = [12, 12, 12, 18, 12, 12, 12]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="KPI_QCP_{date.today()}.xlsx"'
    wb.save(response)
    return response


@login_required
def kpi_comparison(request):
    kpis = get_zone1_kpis()
    return JsonResponse({
        'labels':    [k['name'] for k in kpis],
        'values':    [k['value'] for k in kpis],
        'statuses':  [k['status'] for k in kpis],
        'threshold': 1.0,
        'warning':   0.5,
    })